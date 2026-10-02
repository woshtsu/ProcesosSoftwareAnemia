#!/usr/bin/env python
"""Genera la presentación PPTX (exactamente 7 diapositivas) a partir de la fuente Marp.

Entrada : entregables/semana-06-integrador/presentacion/Presentacion_Integrador.md
Salida  : exportados/semana-06/Presentacion_Integrador_Anemia_Junin.pptx

Uso (desde la raíz del repositorio, con el venv de herramientas):

    herramientas/.venv/Scripts/python entregables/semana-06-integrador/presentacion/generar_presentacion.py
    (Linux/macOS: herramientas/.venv/bin/python ...)

Primera vez (requiere red; solo se instala python-pptx):

    python -m venv herramientas/.venv
    herramientas/.venv/Scripts/python -m pip install python-pptx

El script lee de la fuente Marp: título (# ...), imágenes (![h:N](ruta)), líneas de texto
y notas del orador (comentarios HTML). Inserta las imágenes que existan; si falta alguna,
dibuja un recuadro "imagen pendiente: <ruta>". Al final verifica que haya 7 diapositivas.

Alternativa para PDF/PPTX con Marp (requiere Node y DESCARGAR @marp-team/marp-cli; no la
ejecuta este script, solo se documenta):

    cd entregables/semana-06-integrador/presentacion
    npx @marp-team/marp-cli Presentacion_Integrador.md --pdf --allow-local-files \
        -o ../../../exportados/semana-06/Presentacion_Integrador_Anemia_Junin.pdf
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

AQUI = Path(__file__).resolve().parent
RAIZ = AQUI.parents[2]
FUENTE = AQUI / "Presentacion_Integrador.md"
SALIDA = RAIZ / "exportados" / "semana-06" / "Presentacion_Integrador_Anemia_Junin.pptx"

ANCHO_IN, ALTO_IN = 13.333, 7.5
PX = ANCHO_IN / 1280.0  # pulgadas por píxel Marp (1280x720)
TEAL = RGBColor(0x0B, 0x5C, 0x75)
NARANJA = RGBColor(0xE0, 0x7A, 0x1F)
TINTA = RGBColor(0x1B, 0x2A, 0x3A)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)
GRIS = RGBColor(0x5B, 0x6B, 0x7A)

RE_IMG = re.compile(r"!\[[^\]]*\]\(([^)]+)\)")
RE_ALTO = re.compile(r"!\[[^\]]*?h:(\d+)[^\]]*\]\(([^)]+)\)")
RE_NOTA = re.compile(r"<!--(.*?)-->", re.S)


def leer_diapositivas() -> list[dict]:
    texto = FUENTE.read_text(encoding="utf-8")
    m = re.match(r"^---\s*\n.*?\n---\s*\n", texto, re.S)
    cuerpo = texto[m.end():] if m else texto
    bloques = re.split(r"^---\s*$", cuerpo, flags=re.M)
    diapos = []
    for bloque in bloques:
        if not bloque.strip():
            continue
        notas = [n.strip() for n in RE_NOTA.findall(bloque)]
        limpio = RE_NOTA.sub("", bloque)
        titulo, imagenes, lineas = "", [], []
        for linea in limpio.splitlines():
            s = linea.strip()
            if not s:
                continue
            if s.startswith("# "):
                titulo = s[2:].strip()
            elif RE_IMG.search(s) and not RE_IMG.sub("", s).strip():
                for img in re.finditer(r"!\[([^\]]*)\]\(([^)]+)\)", s):
                    alt, ruta = img.groups()
                    h = re.search(r"h:(\d+)", alt)
                    imagenes.append((ruta, int(h.group(1)) if h else 400))
            else:
                lineas.append(s)
        diapos.append({"titulo": titulo, "imagenes": imagenes, "lineas": lineas,
                       "notas": "\n\n".join(notas)})
    return diapos


def agregar_runs(parrafo, texto: str, tam: int) -> None:
    """Convierte **negrita**, `código` y [COMPLETAR] en runs con formato."""
    texto = texto.replace("`", "")
    partes = re.split(r"(\*\*.+?\*\*|\[COMPLETAR[^\]]*\])", texto)
    for parte in partes:
        if not parte:
            continue
        negrita = parte.startswith("**") and parte.endswith("**")
        completar = parte.startswith("[COMPLETAR")
        run = parrafo.add_run()
        run.text = parte[2:-2] if negrita else parte
        run.font.size = Pt(tam)
        run.font.name = "Calibri"
        run.font.bold = negrita or completar
        run.font.color.rgb = NARANJA if completar else (TEAL if negrita else TINTA)


def resolver(ruta_md: str) -> Path:
    return (AQUI / ruta_md).resolve()


def poner_imagenes(slide, imagenes, region) -> list[str]:
    """Coloca las imágenes en una fila centrada; devuelve las rutas pendientes."""
    x0, y0, ancho, alto = region
    datos = []
    for ruta, h in imagenes:
        p = resolver(ruta)
        if p.exists():
            with Image.open(p) as im:
                asp = im.width / im.height
        else:
            asp = 16 / 9
        datos.append((ruta, p, h * PX, asp))
    hueco = 0.12
    anchos = [h * a for _, _, h, a in datos]
    fila = sum(anchos) + hueco * (len(datos) - 1)
    mayor = max(h for _, _, h, _ in datos)
    esc = min(ancho / fila, alto / mayor, 1.2)
    fila *= esc
    x = x0 + (ancho - fila) / 2
    pendientes = []
    for (ruta, p, h, asp) in datos:
        hh, ww = h * esc, h * asp * esc
        y = y0 + (alto - hh) / 2
        if p.exists():
            slide.shapes.add_picture(str(p), Inches(x), Inches(y), Inches(ww), Inches(hh))
        else:
            rel = p.relative_to(RAIZ).as_posix() if RAIZ in p.parents else ruta
            caja = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
                                          Inches(ww), Inches(hh))
            caja.fill.background()
            caja.line.color.rgb = NARANJA
            caja.line.width = Pt(2)
            caja.line.dash_style = MSO_LINE.DASH
            tf = caja.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            par = tf.paragraphs[0]
            par.alignment = PP_ALIGN.CENTER
            r = par.add_run()
            r.text = f"imagen pendiente: {rel}"
            r.font.size = Pt(18)
            r.font.color.rgb = NARANJA
            pendientes.append(rel)
        x += ww + hueco * esc
    return pendientes


def construir() -> tuple[Presentation, list[str]]:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(ANCHO_IN), Inches(ALTO_IN)
    vacio = prs.slide_layouts[6]
    pendientes: list[str] = []
    for i, d in enumerate(leer_diapositivas(), start=1):
        s = prs.slides.add_slide(vacio)
        # Barra de título
        barra = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.0))
        barra.fill.solid()
        barra.fill.fore_color.rgb = TEAL
        barra.line.fill.background()
        filete = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(1.0), prs.slide_width,
                                    Inches(0.06))
        filete.fill.solid()
        filete.fill.fore_color.rgb = NARANJA
        filete.line.fill.background()
        tf = barra.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = Inches(0.45)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        r = p.add_run()
        r.text = d["titulo"]
        r.font.size = Pt(32)
        r.font.bold = True
        r.font.name = "Calibri"
        r.font.color.rgb = BLANCO
        # Franja de texto inferior (≥ 20 pt)
        n = len(d["lineas"])
        alto_txt = 0.2 + 0.45 * n
        y_txt = ALTO_IN - alto_txt - 0.15
        if n:
            cuadro = s.shapes.add_textbox(Inches(0.4), Inches(y_txt), Inches(ANCHO_IN - 0.8),
                                          Inches(alto_txt))
            t = cuadro.text_frame
            t.word_wrap = True
            t.vertical_anchor = MSO_ANCHOR.MIDDLE
            for k, linea in enumerate(d["lineas"]):
                par = t.paragraphs[0] if k == 0 else t.add_paragraph()
                par.alignment = PP_ALIGN.CENTER
                agregar_runs(par, linea, 20)
        # Imágenes
        if d["imagenes"]:
            region = (0.4, 1.2, ANCHO_IN - 0.8, y_txt - 1.2 - 0.1)
            pendientes += poner_imagenes(s, d["imagenes"], region)
        # Número de diapositiva
        num = s.shapes.add_textbox(Inches(ANCHO_IN - 0.9), Inches(ALTO_IN - 0.4), Inches(0.7),
                                   Inches(0.35))
        pn = num.text_frame.paragraphs[0]
        pn.alignment = PP_ALIGN.RIGHT
        rn = pn.add_run()
        rn.text = f"{i}/7"
        rn.font.size = Pt(18)
        rn.font.color.rgb = GRIS
        # Notas del orador
        s.notes_slide.notes_text_frame.text = d["notas"]
    return prs, pendientes


def main() -> int:
    prs, pendientes = construir()
    SALIDA.parent.mkdir(parents=True, exist_ok=True)
    prs.save(SALIDA)
    comprobada = Presentation(SALIDA)
    n = len(comprobada.slides)
    print(f"PPTX generado: {SALIDA}")
    print(f"Diapositivas: {n}")
    for i, sl in enumerate(comprobada.slides, 1):
        notas = sl.notes_slide.notes_text_frame.text
        print(f"  {i}: notas={len(notas)} caracteres")
    if pendientes:
        print("Imágenes pendientes (recuadro en el PPTX):")
        for ruta in pendientes:
            print(f"  - {ruta}")
    if n != 7:
        print("ERROR: la presentación debe tener exactamente 7 diapositivas", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
